import io
from pptx import Presentation

def extract_text_from_pptx(pptx_bytes: bytes) -> dict:
    prs = Presentation(io.BytesIO(pptx_bytes))
    pages = []

    for i, slide in enumerate(prs.slides, start=1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
        pages.append({
            "page_number": i,
            "text": "\n\n".join(texts),
            "meta": {}
        })

    return {"pages": pages, "meta": {"slides": len(pages)}}
